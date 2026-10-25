"""Convert `ARG27 - Goals, Milestones, Tasks, and Timeline.xlsx` into
src/data/seed.json.

Run from the repo root:  python arg27-tracker/scripts/build-seed.py

Every subteam is imported. Each one has its own `<CODE> Tasks` page (the
`Task Master` sheet only aggregates them) plus its rows on the shared Goals and
Milestones sheets; the app shows one subteam at a time. Nothing is invented:
blank cells stay blank so the app's Data Health view can report on them honestly.
"""

import datetime
import json
import os
import re

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
XLSX = os.path.join(ROOT, "ARG27 - Goals, Milestones, Tasks, and Timeline.xlsx")
OUT = os.path.join(os.path.dirname(HERE), "src", "data", "seed.json")

# The sheet's dropdowns drifted over time; fold the variants onto one vocabulary.
STATUS_MAP = {
    "not started": "Not Started",
    "planned": "Planned",
    "in progress": "In Progress",
    "in review": "In Review",
    "review": "In Review",
    "blocked": "Blocked",
    "complete": "Complete",
    "completed": "Complete",
    "done": "Complete",
    "on track": "In Progress",
    "delayed": "Blocked",
}
GOAL_STATUS_MAP = {
    "on track": "On Track",
    "at risk": "At Risk",
    "delayed": "Delayed",
    "complete": "Complete",
    "done": "Complete",
}
PRIORITY_MAP = {
    "critical": "Critical",
    "high": "High",
    "medium": "Medium",
    "low": "Low",
}

ID_RE = re.compile(r"^[A-Z]{2}(_[A-Z]?\d+)*_?[A-Z]?\d{3}$")

# Assemblies the team has since combined into one role. The workbook still
# lists them apart -- it is edited by hand all season, and renaming a
# `Part Design` block there would orphan every ID under it -- so the fold
# happens on the way in. `src/lib/merges.ts` carries the same list for the
# in-app importer; the two must stay in step or a re-import would silently
# change the data.
ASSEMBLY_MERGES = [
    {
        "subteam": "Autonomy",
        "from": ["A05"],
        "into": {"code": "A04", "name": "Sensors, Harnessing, and Compute"},
    },
]


def cell(v):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, datetime.date):
        return v.strftime("%Y-%m-%d")
    return re.sub(r"[ \t]+", " ", str(v)).strip()


def day(value):
    """A calendar day, or blank. One Due cell on the Milestones sheet holds the
    fragment `8/`; a half-typed date is not a date, so it imports empty and the
    app's Data Health view reports the milestone as undated."""
    return value if re.match(r"^\d{4}-\d{2}-\d{2}$", value or "") else ""


def read(ws, header_row=0):
    """Rows as dicts. The per-subteam task pages repeat their header above every
    assembly block, so any row that restates the header is dropped."""
    rows = [[cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
    if not rows:
        return []
    headers = rows[header_row]
    out = []
    for raw in rows[header_row + 1:]:
        if raw == headers:
            continue
        rec = {h: v for h, v in zip(headers, raw) if h}
        if any(rec.values()):
            out.append(rec)
    return out


def header_row_of(ws):
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if cell(row[0]) == "ID":
            return i
    return 0


def split_owners(value):
    """One or more people per cell, separated by punctuation or `and`."""
    parts = re.split(r"\s*(?:,|;|/|&|\band\b)\s*", value or "")
    return [p.strip() for p in parts if p.strip()]


def split_refs(value):
    """`Dependencies` cells hold one or more IDs, comma/space separated."""
    return [p for p in re.split(r"[,;/]|\s+", value or "") if ID_RE.match(p)]


def split_assembly(label):
    """`A06 - Perception and Data Interface` -> ("A06", "Perception ...")."""
    m = re.match(r"^([A-Z]\d+)\s*-\s*(.+)$", label or "")
    return (m.group(1), m.group(2)) if m else ("", label or "")


def month_floor(day):
    return day[:7] + "-01" if day else ""


def month_ceil(day):
    if not day:
        return ""
    y, m = int(day[:4]), int(day[5:7])
    y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return "%04d-%02d-01" % (y, m)


def apply_assembly_merges(subteams, assemblies, tasks, milestones):
    """Fold the merged assemblies together, in place.

    Tasks moving across keep their place in sheet order and are renumbered
    into the surviving assembly's sequence -- `AU_A05_T003` becomes
    `AU_A04_T005` -- because the ID carries the assembly code everywhere the
    app shows it. The workbook's own ID stays on the task as `sheetId`, and
    references to it from other tasks and milestones follow the renumber.

    Run before IDs are deduplicated, so a renumber that lands on an ID the
    sheet already used is caught by that pass like any other collision."""
    for merge in ASSEMBLY_MERGES:
        sub = next((s for s in subteams if s["name"] == merge["subteam"]), None)
        if not sub:
            continue
        into = merge["into"]
        codes = set([into["code"]] + merge["from"])

        # One assembly where there were several, in the position of the first.
        kept, keep = False, []
        for a in assemblies:
            if a["subteam"] != merge["subteam"] or a["code"] not in codes:
                keep.append(a)
                continue
            if kept:
                continue
            kept = True
            a["code"], a["name"] = into["code"], into["name"]
            keep.append(a)
        assemblies[:] = keep

        renamed = {}
        moved = [t for t in tasks
                 if t["subteam"] == merge["subteam"] and t["assemblyCode"] in codes]
        for i, t in enumerate(moved, 1):
            t["assembly"], t["assemblyCode"] = into["name"], into["code"]
            new = "%s_%s_T%03d" % (sub["code"], into["code"], i)
            if new != t["id"]:
                if t["id"]:
                    renamed[t["id"]] = new
                    t.setdefault("sheetId", t["id"])
                t["id"] = new

        if not renamed:
            continue
        for it in tasks + milestones:
            it["dependsOn"] = [renamed.get(x, x) for x in it["dependsOn"]]


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    # The References sheet is the authority on which subteams exist and what
    # each one is abbreviated to.
    subteams = []
    for r in read(wb["References"]):
        name, code = r.get("Subteam", ""), r.get("ID", "")
        if name and code and not any(s["code"] == code for s in subteams):
            subteams.append({"code": code, "name": name})

    goals, milestones, tasks, assemblies = [], [], [], []

    goals_raw = read(wb["Goals"])
    ms_raw = read(wb["Milestones"])

    for st in subteams:
        name, code = st["name"], st["code"]

        # Goals. The Goals sheet numbers them in the milestone namespace
        # (`AU_M001` is both a goal and a milestone), so re-key them here and
        # keep the sheet's own ID as a breadcrumb.
        n = 0
        for r in goals_raw:
            if r.get("Subteam") != name or not r.get("Goal"):
                continue
            n += 1
            goals.append(
                {
                    "id": "%s_G%03d" % (code, n),
                    "sheetId": r.get("ID", ""),
                    "subteam": name,
                    "title": r.get("Goal", ""),
                    "status": GOAL_STATUS_MAP.get(r.get("Status", "").lower(), ""),
                    "successMetric": r.get("Success Metric", ""),
                    "notes": r.get("Notes", ""),
                }
            )

        m = 0
        for r in ms_raw:
            if r.get("Subteam") != name or not r.get("Task"):
                continue
            m += 1
            milestones.append(
                {
                    "id": r.get("ID", "") or "%s_M%03d" % (code, m),
                    "subteam": name,
                    "title": r.get("Task", ""),
                    "start": day(r.get("Start", "")),
                    "due": day(r.get("Due", "")),
                    "status": STATUS_MAP.get(r.get("Status", "").lower(), ""),
                    # The sheet carries no goal column on milestones; that link
                    # is made in the app and reported as a gap until it is.
                    "goalIds": [],
                    "dependsOn": split_refs(r.get("Dependencies", "")),
                    "notes": r.get("Notes", ""),
                }
            )

        # Tasks live on one page per subteam. A couple of subteams have goals
        # and milestones but no task page at all.
        sheet = "%s Tasks" % code
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        rows = read(ws, header_row_of(ws))

        # Assembly order comes from the task page itself: every row carries its
        # `Part Design` block label, including the blank placeholder rows.
        seen = []
        for r in rows:
            label = r.get("Part Design", "")
            if label and label not in seen:
                seen.append(label)
        for label in seen:
            a_code, a_name = split_assembly(label)
            if a_name:
                assemblies.append({"subteam": name, "code": a_code, "name": a_name})

        for r in rows:
            # Rows carrying only a block label are empty template rows.
            if not r.get("Task"):
                continue
            a_code, a_name = split_assembly(r.get("Part Design", ""))
            tasks.append(
                {
                    "id": r.get("ID", ""),
                    "subteam": name,
                    "title": r.get("Task", ""),
                    "assembly": a_name,
                    "assemblyCode": a_code,
                    "start": day(r.get("Start", "")),
                    "due": day(r.get("Due", "")),
                    "status": STATUS_MAP.get(r.get("Status", "").lower(), ""),
                    "milestoneId": r.get("Milestone", ""),
                    "owners": split_owners(r.get("Owner", "")),
                    "priority": PRIORITY_MAP.get(r.get("Priority", "").lower(), ""),
                    "dependsOn": split_refs(r.get("Dependencies", "")),
                    "deliverable": r.get("Deliverables", ""),
                    "notes": r.get("Notes", ""),
                }
            )

    # Assemblies the team has since combined into one role; the sheet still
    # lists them apart. Folded before the duplicate pass, so a renumbered ID
    # that collides is handled there like any other.
    apply_assembly_merges(subteams, assemblies, tasks, milestones)

    # The sheet reuses auto-generated IDs here and there. Keep such rows
    # distinguishable in the app, with a breadcrumb back to the original.
    for items in (goals, milestones, tasks):
        counts = {}
        for it in items:
            k = it["id"]
            c = counts.get(k, 0)
            counts[k] = c + 1
            if c:
                it["duplicateOf"] = k
                it["id"] = "%s-%d" % (k, c + 1)

    days = sorted(d for i in tasks + milestones for d in (i.get("start"), i.get("due")) if d)
    competition = "2027-06-16"
    data = {
        "meta": {
            "source": os.path.basename(XLSX),
            "generated": True,
            "seasonStart": min(month_floor(days[0]), "2026-08-01") if days else "2026-08-01",
            "seasonEnd": max(month_ceil(days[-1]), month_ceil(competition)) if days else "2027-07-01",
            "competition": competition,
        },
        # Only subteams that actually carry work are worth switching to.
        "subteams": [
            s
            for s in subteams
            if any(i["subteam"] == s["name"] for i in tasks + milestones + goals)
        ],
        "assemblies": assemblies,
        "goals": goals,
        "milestones": milestones,
        "tasks": tasks,
    }
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as fh:
        json.dump(data, fh, indent=1, ensure_ascii=False)

    print("subteams=%d assemblies=%d goals=%d milestones=%d tasks=%d -> %s" % (
        len(data["subteams"]), len(assemblies), len(goals), len(milestones), len(tasks), OUT))
    for s in data["subteams"]:
        print("  %-14s %3d tasks  %2d milestones  %d goals" % (
            s["name"],
            sum(1 for t in tasks if t["subteam"] == s["name"]),
            sum(1 for x in milestones if x["subteam"] == s["name"]),
            sum(1 for g in goals if g["subteam"] == s["name"]),
        ))


if __name__ == "__main__":
    main()
